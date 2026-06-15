from __future__ import annotations

import json
import math
import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl


APP_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = APP_ROOT / "src" / "data" / "generated"
PUBLIC_ASSETS = APP_ROOT / "public" / "generated" / "assets"
SOURCE_NAME = "【可任意分享】健身Excel超级套表（作者：B站好人松松）26年4月最新版.xlsx"


MODULES = [
    {"id": "dashboard", "label": "总览", "description": "按目标学习整本健身 Excel。"},
    {"id": "fat-loss", "label": "减脂饮食", "description": "不同训练时间的减脂饮食方案。"},
    {"id": "muscle-gain", "label": "增肌饮食", "description": "不同训练时间的增肌饮食方案。"},
    {"id": "training", "label": "训练计划", "description": "健身房和居家分化训练安排。"},
    {"id": "tools", "label": "热量工具", "description": "有氧消耗、1RM 和食物重量换算。"},
    {"id": "foods", "label": "食物营养", "description": "日常碳水和蛋白质营养率。"},
    {"id": "qa", "label": "问答库", "description": "减脂与增肌执行问题。"},
    {"id": "anatomy", "label": "拉伸解剖", "description": "拉伸图谱、关节活动与肌肉关系。"},
    {"id": "source", "label": "原表索引", "description": "30 个 sheet 的分类地图和原始行。"},
]


def module_for_sheet(title: str) -> str:
    if "减脂-" in title and "问答" not in title:
        return "fat-loss"
    if "增肌-" in title and "问答" not in title:
        return "muscle-gain"
    if "训练计划" in title:
        return "training"
    if "有氧" in title or "最大力量" in title:
        return "tools"
    if "食物" in title:
        return "foods"
    if "问答" in title:
        return "qa"
    if any(token in title for token in ["拉伸", "解剖", "关节活动", "肌肉的关节"]):
        return "anatomy"
    return "dashboard"


def display_title(title: str) -> str:
    return re.sub(r"^\d+", "", title).strip("-")


def source_ref(ws, row: int | None = None, row_range: str | None = None) -> dict:
    ref = {
        "sheetIndex": ws.parent.worksheets.index(ws) + 1,
        "sheetTitle": ws.title,
    }
    if row is not None:
        ref["row"] = row
    if row_range is not None:
        ref["rowRange"] = row_range
    return ref


def workbook_source() -> Path:
    preferred = PUBLIC_ASSETS / SOURCE_NAME
    if preferred.exists():
        return preferred
    matches = sorted(PUBLIC_ASSETS.glob("*.xlsx"))
    if len(matches) == 1:
        return matches[0]
    raise FileNotFoundError(f"Expected one workbook in {PUBLIC_ASSETS}, found {len(matches)}")


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def compact(value: str, limit: int = 180) -> str:
    value = re.sub(r"\s+", " ", clean(value).replace("\n", " "))
    return value if len(value) <= limit else value[: limit - 1] + "…"


def slugify(text: str) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "-", text).strip("-")


def image_ext(data: bytes) -> str:
    if data.startswith(b"\x89PNG"):
        return ".png"
    if data.startswith(b"\xff\xd8"):
        return ".jpg"
    if data.startswith(b"GIF8"):
        return ".gif"
    return ".bin"


def image_anchor(img) -> tuple[int, int]:
    marker = getattr(getattr(img, "anchor", None), "_from", None)
    if marker is None:
        return (9999, 9999)
    return (int(marker.row), int(marker.col))


def nonempty_rows(ws, max_cols: int | None = None) -> list[dict]:
    rows = []
    col_limit = max_cols or ws.max_column
    for r, row in enumerate(ws.iter_rows(values_only=True), 1):
        values = [clean(v) for v in row[:col_limit]]
        if any(values):
            rows.append({"row": r, "values": values})
    return rows


def row_values(ws, row: int, cols: int | None = None) -> list[str]:
    limit = cols or ws.max_column
    return [clean(ws.cell(row, c).value) for c in range(1, limit + 1)]


MEAL_PROFILES = {
    "breakfast": {"label": "早饭", "carbPercent": 20, "proteinPercent": 20},
    "preworkout": {"label": "练前餐", "carbPercent": 15, "proteinPercent": 0},
    "postworkout": {"label": "练后餐", "carbPercent": 35, "proteinPercent": 30},
    "regular": {"label": "其他餐", "carbPercent": 20, "proteinPercent": 30},
    "snack": {"label": "零食/夜宵", "carbPercent": 10, "proteinPercent": 20},
    "rest-main": {"label": "正餐", "carbPercent": 35, "proteinPercent": 30},
}


def meal_type_for_name(name: str, day_type: str) -> str:
    if "零食" in name or "夜宵" in name:
        return "snack"
    if "练前" in name:
        return "preworkout"
    if "练后" in name:
        return "postworkout"
    if "早饭" in name:
        return "breakfast"
    if day_type in {"rest", "daily"} and ("午饭" in name or "晚饭" in name):
        return "rest-main"
    return "regular"


def collect_meal_notes(ws, start_row: int, end_row: int, max_cols: int = 19) -> dict:
    carb_options = []
    protein_options = []
    fat_notes = []
    produce_notes = []
    for row in range(start_row + 1, end_row):
        values = row_values(ws, row, max_cols)
        for value in values[4:8]:
            if value and value != "0" and not re.fullmatch(r"\d+\.?", value):
                carb_options.append(compact(value, 150))
        for value in values[8:12]:
            if value and value != "0" and not re.fullmatch(r"\d+\.?", value):
                protein_options.append(compact(value, 160))
        for value in values[12:13]:
            if value and value != "0":
                fat_notes.append(compact(value, 220))
        for value in values[13:19]:
            if value and value != "0":
                produce_notes.append(compact(value, 220))
    return {
        "carbOptions": carb_options[:8],
        "proteinOptions": protein_options[:8],
        "fatNote": " ".join(fat_notes[:3]),
        "produceNote": " ".join(produce_notes[:3]),
    }


def parse_meal_table(ws, day_type: str, title: str, start_row: int, end_row: int) -> dict:
    meal_rows = []
    for row in range(start_row, end_row + 1):
        name = clean(ws.cell(row, 2).value)
        if re.match(r"^[①②③④⑤]", name):
            meal_rows.append((row, name))

    meals = []
    for index, (row, name) in enumerate(meal_rows):
        next_row = meal_rows[index + 1][0] if index + 1 < len(meal_rows) else end_row + 1
        meal_type = meal_type_for_name(name, day_type)
        profile = MEAL_PROFILES[meal_type]
        notes = collect_meal_notes(ws, row, next_row)
        meals.append({
            "order": re.match(r"^[①②③④⑤]", name).group(0),
            "name": name,
            "mealType": meal_type,
            "mealTypeLabel": profile["label"],
            "carbPercent": profile["carbPercent"],
            "proteinPercent": profile["proteinPercent"],
            "carbOptions": notes["carbOptions"],
            "proteinOptions": ["蛋白质不用吃"] if meal_type == "preworkout" else notes["proteinOptions"],
            "fatNote": notes["fatNote"],
            "produceNote": notes["produceNote"],
            "sourceRefs": [source_ref(ws, row_range=f"{row}:{next_row - 1}")],
        })
    return {
        "dayType": day_type,
        "title": title,
        "meals": meals,
        "sourceRefs": [source_ref(ws, row_range=f"{start_row}:{end_row}")],
    }


def parse_diet_meal_tables(ws, goal: str) -> list[dict]:
    if goal != "fat-loss":
        return []
    tables = []
    markers = []
    for row in range(1, ws.max_row + 1):
        text = " ".join(row_values(ws, row, min(ws.max_column, 19)))
        if "力训日饮食" in text:
            markers.append(("training", "力训日饮食", row))
        elif "休息日饮食" in text:
            markers.append(("rest", "休息日饮食", row))
        elif "每日饮食" in text:
            markers.append(("daily", "每日饮食", row))
    for index, (day_type, title, marker_row) in enumerate(markers):
        next_marker = markers[index + 1][2] if index + 1 < len(markers) else ws.max_row + 1
        end_row = next_marker - 1
        for row in range(marker_row + 1, next_marker):
            text = " ".join(row_values(ws, row, min(ws.max_column, 19)))
            if "训练计划" in text or "有氧方案" in text:
                end_row = row - 1
                break
        table = parse_meal_table(ws, day_type, title, marker_row, end_row)
        if table["meals"]:
            tables.append(table)
    return tables


def save_images(ws) -> list[dict]:
    assets = []
    for idx, img in enumerate(sorted(getattr(ws, "_images", []), key=image_anchor), 1):
        data = img._data()
        ext = image_ext(data)
        filename = f"{slugify(ws.title)}-{idx:02d}{ext}"
        path = PUBLIC_ASSETS / filename
        path.write_bytes(data)
        row, col = image_anchor(img)
        assets.append(
            {
                "sheet": ws.title,
                "title": ws.title,
                "src": f"generated/assets/{filename}",
                "label": f"图 {idx}",
                "anchor": f"约 Excel 第 {row + 1} 行 / 第 {col + 1} 列",
                "sourceRefs": [source_ref(ws, row + 1)],
            }
        )
    return assets


def parse_diet_plan(ws, idx: int) -> dict:
    title = ws.title
    goal = "fat-loss" if "减脂" in title else "muscle-gain"
    timing = re.sub(r"^\d+(减脂|增肌)-", "", title)
    rows = nonempty_rows(ws, 12)
    notes = []
    inputs = []
    training_day = []
    rest_day = []
    daily_meals = []
    current_section = "summary"
    for item in rows:
        text = " ".join(v for v in item["values"] if v)
        if not text or text in {"HR", "H", "R", "S"}:
            continue
        if "力训日饮食" in text:
            current_section = "training"
            continue
        if "休息日饮食" in text:
            current_section = "rest"
            continue
        if "每日饮食" in text:
            current_section = "daily"
            continue
        if item["row"] <= 28:
            label = next((v for v in item["values"][1:5] if v and v not in {"自己输入"}), "")
            detail = " ".join(v for v in item["values"][2:] if v and v != label)
            if label and len(label) < 24:
                inputs.append({"label": label, "detail": compact(detail, 260)})
            else:
                notes.append(compact(text, 300))
        elif current_section == "training":
            training_day.append({"row": item["row"], "text": compact(text, 320), "sourceRefs": [source_ref(ws, item["row"])]})
        elif current_section == "rest":
            rest_day.append({"row": item["row"], "text": compact(text, 320), "sourceRefs": [source_ref(ws, item["row"])]})
        elif current_section == "daily":
            daily_meals.append({"row": item["row"], "text": compact(text, 320), "sourceRefs": [source_ref(ws, item["row"])]})
        else:
            notes.append(compact(text, 260))
    summary = "；".join(n for n in notes[:3] if n)
    meal_tables = parse_diet_meal_tables(ws, goal)
    return {
        "id": slugify(title),
        "sheetIndex": idx,
        "title": title,
        "goal": goal,
        "timing": timing,
        "summary": summary,
        "inputs": inputs[:12],
        "trainingDayMeals": training_day[:80],
        "restDayMeals": rest_day[:80],
        "dailyMeals": daily_meals[:80],
        "mealTables": meal_tables,
        "notes": notes[:12],
        "rawRows": rows,
        "sourceRefs": [source_ref(ws, row_range=f"1:{ws.max_row}")],
    }


def parse_cardio(ws) -> list[dict]:
    headers = row_values(ws, 12, 21)
    weights = []
    for i, value in enumerate(headers, 1):
        if value.isdigit():
            weights.append((i, value))
    entries = []
    current_rest = ""
    for r in range(14, ws.max_row + 1):
        values = row_values(ws, r, 21)
        row_text = " ".join(values)
        rest_match = re.search(r"静息心率(\d+)", row_text)
        if rest_match:
            current_rest = rest_match.group(1)
        exercise_match = re.search(r"运动心率(\d+)", row_text)
        kcal = values[4] if len(values) > 4 else ""
        if exercise_match and kcal:
            weight_map = {weight: values[col - 1] for col, weight in weights if col - 1 < len(values) and values[col - 1]}
            entries.append(
                {
                    "method": "heart-rate",
                    "restingHeartRate": current_rest,
                    "exerciseHeartRate": exercise_match.group(1),
                    "kcalPerKg": float(kcal) if re.match(r"^\d+(\.\d+)?$", kcal) else kcal,
                    "weightKcalMap": weight_map,
                    "sourceRefs": [source_ref(ws, r)],
                }
            )
    return entries


def parse_qa(ws, category: str) -> list[dict]:
    rows = nonempty_rows(ws, 3)
    items = []
    current = None
    for item in rows:
        text = " ".join(v for v in item["values"] if v)
        if not text or text in {"H H", "R R", "S S"}:
            continue
        match = re.search(r"(\d+)[\.．]\s*([^？?]+[？?]?)", text)
        is_directory = item["row"] < 23
        if match and not is_directory:
            if current:
                items.append(current)
            question = match.group(2).strip()
            current = {
                "id": f"{category}-{match.group(1)}",
                "category": category,
                "question": question,
                "answer": "",
                "tags": list({category, *re.findall(r"减脂|增肌|有氧|外卖|食堂|体重|体脂|痛风|糖尿病|练前|练后", question)}),
                "sourceRefs": [source_ref(ws, item["row"])],
            }
            remainder = text[match.end() :].strip()
            if remainder:
                current["answer"] += remainder
        elif current:
            current["answer"] += ("\n" if current["answer"] else "") + text
    if current:
        items.append(current)

    if not items:
        directory = [item for item in rows if 4 <= item["row"] <= 29]
        for item in directory:
            text = " ".join(v for v in item["values"] if v)
            match = re.search(r"(\d+)[\.．]\s*(.+)", text)
            if match:
                items.append(
                    {
                        "id": f"{category}-{match.group(1)}",
                        "category": category,
                        "question": match.group(2),
                        "answer": "原表中该问题为目录条目，详情请查看原文行。",
                        "tags": [category],
                        "sourceRefs": [source_ref(ws, item["row"])],
                    }
                )
    return items


def parse_foods(ws) -> list[dict]:
    foods = []
    macro = ""
    current_group = ""
    for r in range(1, ws.max_row + 1):
        values = row_values(ws, r, 6)
        if values[1] in {"碳水", "蛋白质"}:
            macro = values[1]
            current_group = ""
            continue
        if values[1] == "大类":
            continue
        if values[1]:
            current_group = values[1]
        name = values[2]
        rate = values[3]
        if macro and name and rate:
            foods.append(
                {
                    "id": f"{slugify(macro)}-{r}",
                    "macroType": macro,
                    "group": current_group,
                    "name": name,
                    "rate": rate,
                    "giOrPosition": values[4],
                    "explanation": values[5],
                    "sourceRefs": [source_ref(ws, r)],
                }
            )
    return foods


def parse_training_plan(ws, idx: int) -> dict:
    rows = nonempty_rows(ws, 7)
    title = next((r["values"][1] for r in rows if len(r["values"]) > 1 and r["values"][1]), ws.title)
    info = []
    days = []
    current = None
    current_group = ""
    current_sets = ""
    for item in rows:
        values = item["values"]
        text = " ".join(v for v in values if v)
        if item["row"] < 16 and values[1] and values[2]:
            info.append({"label": values[1], "detail": values[2]})
        day_title = next((v for v in values if v.startswith("Day")), "")
        if day_title:
            if current:
                days.append(current)
            current = {"title": day_title, "rationale": "", "rows": []}
            current_group = ""
            current_sets = ""
            continue
        if not current:
            continue
        if "为什么" in text and not current["rationale"]:
            current["rationale"] = text
            continue
        if values[1] == "部位" or values[0] in {"H", "R", "S"}:
            continue
        if values[1]:
            current_group = values[1]
        if values[2]:
            current_sets = values[2]
        exercise = values[3] if len(values) > 3 else ""
        if exercise or values[4] or values[5]:
            current["rows"].append(
                {
                    "muscleGroup": current_group,
                    "setGuidance": current_sets,
                    "exercise": exercise,
                    "shoulderJoint": values[4] if len(values) > 4 else "",
                    "elbowJoint": values[5] if len(values) > 5 else "",
                    "notes": values[6] if len(values) > 6 else "",
                    "sourceRefs": [source_ref(ws, item["row"])],
                }
            )
    if current:
        days.append(current)
    return {
        "id": slugify(ws.title),
        "sheetIndex": idx,
        "title": title,
        "environment": "home" if "居家" in ws.title else "gym",
        "splitType": "四分化" if "四分化" in ws.title else "三分化",
        "info": info,
        "days": days,
        "rawRows": rows,
        "sourceRefs": [source_ref(ws, row_range=f"1:{ws.max_row}")],
    }


def parse_one_rep_max(ws) -> list[dict]:
    formulas = [
        ("Adams", "w/(1-0.02*r)", "对自由卧推/深蹲做组估算"),
        ("Brown", "(r*0.0328+0.9849)*w", "对女性较准确"),
        ("Brzycki", "w/(1.0278-0.0278*r)", "对女性较准确"),
        ("Lander", "w/(1.013-0.0267123*r)", "对女性较准确"),
        ("Lombardi", "w*(r**0.1)", "对男性较准确"),
        ("Mayhew", "w/(0.522+0.419*math.exp(-0.055*r))", ""),
        ("O’Connor", "0.025*w*r+w", ""),
        ("Wathen", "w/(0.488+0.538*math.exp(-0.075*r))", ""),
        ("Welday", "(r*0.0333)*w+w", ""),
    ]
    return [
        {
            "author": author,
            "expression": expr,
            "note": note,
            "example": round(eval(expr, {"math": math}, {"w": 50, "r": 10}), 2),
            "sourceRefs": [source_ref(ws, row)],
        }
        for row, (author, expr, note) in enumerate(formulas, 5)
    ]


def parse_anatomy(wb) -> dict:
    ws = wb["27健身解剖总结（文字版）"]
    intro = []
    for r in (4, 5, 6):
        values = row_values(ws, r, 10)
        intro.append({"title": values[1], "body": values[3]})

    joint_rows = []
    current_joint = ""
    for r in range(17, 35):
        values = row_values(ws, r, 10)
        current_joint = values[1] or current_joint
        joint_rows.append(
            {
                "joint": current_joint,
                "movement": values[2],
                "description": values[3],
                "example": values[4],
                "muscles": [v for v in values[5:10] if v],
                "sourceRefs": [source_ref(ws, r)],
            }
        )

    def muscle_section(title_row, header_row, start, end):
        headers = row_values(ws, header_row, 12)
        muscle_headers = [(i, h) for i, h in enumerate(headers[4:], 4) if h]
        current = ""
        rows = []
        for r in range(start, end + 1):
            values = row_values(ws, r, 12)
            current = values[1] or current
            rows.append(
                {
                    "joint": current,
                    "movement": values[2],
                    "description": values[3],
                    "targets": [
                        {"muscle": h, "item": values[i]}
                        for i, h in muscle_headers
                        if i < len(values) and values[i]
                    ],
                    "sourceRefs": [source_ref(ws, r)],
                }
            )
        return {"title": clean(ws.cell(title_row, 2).value), "rows": rows}

    galleries = []
    for title in ["25上身肌肉拉伸", "26下身肌肉拉伸", "28关节活动的肌肉（图示版）", "29肌肉的关节活动（图示版）"]:
        image_assets = save_images(wb[title])
        note = next((r["values"][0] for r in nonempty_rows(wb[title], 7) if r["values"][0]), "")
        galleries.append({"id": slugify(title), "title": title, "note": note, "images": image_assets})

    return {
        "intro": intro,
        "jointToMuscles": joint_rows,
        "muscleSections": [
            muscle_section(42, 44, 45, 50),
            muscle_section(53, 55, 56, 61),
            muscle_section(63, 65, 66, 67),
            muscle_section(70, 72, 73, 76),
        ],
        "notes": [clean(ws.cell(r, 2).value) for r in (35, 51, 62, 68) if clean(ws.cell(r, 2).value)],
        "imageGalleries": galleries,
    }


def main() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PUBLIC_ASSETS.mkdir(parents=True, exist_ok=True)
    for path in PUBLIC_ASSETS.iterdir():
        if path.suffix.lower() == ".xlsx":
            continue
        if path.is_dir():
            shutil.rmtree(path)
        else:
            path.unlink()

    source = workbook_source()
    wb = openpyxl.load_workbook(source, data_only=True)
    sheets = [{"index": i, "title": ws.title, "rows": ws.max_row, "cols": ws.max_column} for i, ws in enumerate(wb.worksheets, 1)]
    sheet_categories = [
        {
            "index": item["index"],
            "title": item["title"],
            "displayTitle": display_title(item["title"]),
            "moduleId": module_for_sheet(item["title"]),
            "rows": item["rows"],
            "cols": item["cols"],
        }
        for item in sheets
    ]

    diet_plans = [parse_diet_plan(ws, i) for i, ws in enumerate(wb.worksheets[1:16], 2)]
    training = [parse_training_plan(wb[name], i) for i, name in enumerate([
        "20训练计划-健身房三分化",
        "21训练计划-健身房四分化计划（肩单练版）",
        "22训练计划-健身房四分化计划（手臂单练版）",
        "23训练计划-居家健身",
    ], 20)]

    qa = parse_qa(wb["17减脂-问答汇总"], "减脂") + parse_qa(wb["18增肌-问答汇总"], "增肌")
    raw_sheets = [
        {
            "index": i,
            "title": ws.title,
            "rows": nonempty_rows(ws, min(ws.max_column, 12)),
        }
        for i, ws in enumerate(wb.worksheets, 1)
    ]

    data = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "source": source.name,
        "sheets": sheets,
        "modules": MODULES,
        "sheetCategories": sheet_categories,
        "dietPlans": diet_plans,
        "cardio": parse_cardio(wb["16有氧热量消耗"]),
        "qa": qa,
        "foods": parse_foods(wb["19日常食物营养率"]),
        "trainingPlans": training,
        "oneRepMax": parse_one_rep_max(wb["24最大力量预测公式"]),
        "anatomy": parse_anatomy(wb),
        "rawSheets": raw_sheets,
    }
    (DATA_DIR / "workbook.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "sheets": len(sheets),
        "dietPlans": len(diet_plans),
        "qa": len(qa),
        "foods": len(data["foods"]),
        "trainingPlans": len(training),
        "cardio": len(data["cardio"]),
        "images": sum(len(g["images"]) for g in data["anatomy"]["imageGalleries"]),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
